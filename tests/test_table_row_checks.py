"""Original blank/content contradictions and finite review, not model quality."""

import copy
import io
import json
import zipfile
from xml.sax.saxutils import escape

import pytest
from openpyxl import Workbook
from test_native_merged_geometry import observe
from test_table_layout import Model, fixture, mapping, response, run

from document_files.interpretation import table_layout as layout
from document_files.interpretation.backends import InferenceResponse
from document_files.interpretation.compiler import CompileError, compile_region
from document_files.interpretation.regions import prepare_regions
from document_files.interpretation.table_protocol import structural_ir
from document_files.interpretation.table_row_checks import BLANK_ROW_ERROR, blank_row_conflicts
from document_files.interpretation.table_source_wire import expand_table_sources
from document_files.interpretation.table_structure_feedback import (
    needs_layout_review,
    structure_feedback,
)


def native_file(format_id, text):
    stream = io.BytesIO()
    if format_id == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "Heading"
        sheet["A2"] = text
        sheet["A2"].number_format = "@"  # Preserve a physically stored empty XML cell.
        book.save(stream)
        book.close()
    else:
        with zipfile.ZipFile(stream, "w") as out:
            out.writestr(
                "Contents/section0.xml",
                '<sec><p><run><tbl rowCnt="2" colCnt="1">'
                + "".join(
                    f'<tr><tc><cellAddr rowAddr="{row}" colAddr="0"/>'
                    '<cellSpan rowSpan="1" colSpan="1"/><subList><p><run>'
                    f"<t>{escape(value)}</t></run></p></subList></tc></tr>"
                    for row, value in enumerate(["Heading", text])
                )
                + "</tbl></run></p></sec>",
            )
    return stream.getvalue()


@pytest.mark.parametrize("format_id", ["hwpx", "xlsx"])
@pytest.mark.parametrize("text", ["", " ", "0", "false", "=1-1", "Blank means no result.", "주석"])
def test_native_blank_validation_preserves_exact_empty_whitespace_number_formula_and_note(
    format_id, text
):
    doc = observe(native_file(format_id, text), format_id)
    region = prepare_regions(doc, context_chars=16000)[0]
    before = copy.deepcopy(doc)
    value = response(doc, region, ["header", "blank"])
    if text == "":
        accepted = layout.accept(value, doc, region)
        assert accepted["response"] == value
    else:
        with pytest.raises(CompileError, match=BLANK_ROW_ERROR) as caught:
            layout.accept(value, doc, region)
        detail = caught.value.selection["rows"]
        assert [d["row"] for d in detail] == [1]
        assert all(doc.nodes[r] for r in detail[0]["nonemptySourceRefs"])
        feedback = structure_feedback(caught.value, doc, region)
        assert len(feedback) == 2 and feedback[0] == BLANK_ROW_ERROR
        assert all(k not in feedback[1] for k in ['"text"', '"value"', '"role"'])
    assert doc == before


@pytest.mark.parametrize("value", [0, False])
def test_real_xlsx_zero_and_false_are_not_empty(value):
    doc = observe(native_file("xlsx", value), "xlsx")
    region = prepare_regions(doc, context_chars=16000)[0]
    with pytest.raises(CompileError, match=BLANK_ROW_ERROR):
        layout.accept(response(doc, region, ["header", "blank"]), doc, region)


def test_merged_span_and_multiple_paragraph_sources_are_all_inspected_without_using_empty_excerpt():
    doc, region, _ = fixture(
        b'<table><tr><td rowspan="2">Marker</td><td></td></tr><tr><td></td></tr></table>'
    )
    table = doc.tables[region["tableRef"]]
    ref = table["cells"][0]["sourceRef"]
    region["nodeViews"] = {ref: {"start": 0, "end": 0}}
    before = copy.deepcopy(doc)
    conflicts = blank_row_conflicts(doc, table, {0: "blank", 1: "blank"})
    assert [i["row"] for i in conflicts] == [0, 1]
    assert all(i["nonemptySourceRefs"] == [ref] for i in conflicts)
    assert doc == before
    # The second original paragraph is not hidden by an empty cell anchor.
    doc.nodes[ref]["text"] = ""
    doc.nodes["paragraph"] = {"text": "Original second paragraph"}
    table["cells"][0]["sourceRefs"] = [ref, "paragraph"]
    assert all(
        i["nonemptySourceRefs"] == ["paragraph"]
        for i in blank_row_conflicts(doc, table, {0: "blank", 1: "blank"})
    )


@pytest.mark.parametrize(
    "kind", ["binding_conflict", "node_conflict", "unknown_text", "missing_node"]
)
def test_unclear_empty_observation_is_not_reported_as_nonempty_or_accepted_blank(kind):
    doc, region, _ = fixture(b"<table><tr><td></td></tr></table>")
    table = doc.tables[region["tableRef"]]
    ref = table["cells"][0]["sourceRef"]
    if kind == "binding_conflict":
        doc.bindings["conflict"] = {"sourceRef": ref, "candidateStatus": "unresolved_conflict"}
    elif kind == "node_conflict":
        doc.nodes[ref]["semanticInput"] = {"role": "unresolved_conflict"}
    elif kind == "unknown_text":
        doc.nodes[ref]["text"] = None
    else:
        del doc.nodes[ref]
    before = copy.deepcopy(doc)
    with pytest.raises(CompileError, match=BLANK_ROW_ERROR) as caught:
        layout.accept(response(doc, region, ["blank"]), doc, region)
    assert caught.value.selection == {
        "rows": [{"row": 0, "nonemptySourceRefs": [], "unverifiedSourceRefs": [ref]}]
    }
    assert doc == before


def test_compiler_also_rejects_nonempty_blank_rows_on_the_canonical_path():
    doc, region, payload = fixture()
    current = layout.accept(response(doc, region), doc, region)
    value = mapping(expand_table_sources(layout.mapping_request(payload, current, doc, region)))
    _, ir = structural_ir(layout.decode_mapping(value, current, doc, region), doc, region)
    original = copy.deepcopy((doc, ir))
    ir.repeats[0].rowRoles[-1].role = "blank"
    with pytest.raises(CompileError, match=BLANK_ROW_ERROR):
        compile_region(ir, doc, region)
    assert doc == original[0] and ir.repeats[0].rowRoles[-1].role == "blank"


class BlankFirst(Model):
    def __init__(self, *, fix=True):
        super().__init__()
        self.fix = fix
        self.layout_calls = 0

    def infer(self, request):
        response = super().infer(request)
        payload = json.loads(request.messages[-1]["content"])
        if payload["tableStage"] == "layout":
            self.layout_calls += 1
            if self.layout_calls == 1 or not self.fix:
                value = json.loads(response.text)
                value["rowRoles"][-1] = "blank"
                return InferenceResponse(json.dumps(value), response.usage)
        return response


@pytest.mark.parametrize("fix", [False, True])
def test_blank_repair_gets_full_source_and_source_rows_before_any_mapping(fix):
    model, states = BlankFirst(fix=fix), []
    result = run(model, states=states, maxModelCalls=4)
    assert len(model.requests) == (4 if fix else 2)
    assert (result["data"] is not None) == fix
    a, b = [expand_table_sources(json.loads(r.messages[-1]["content"])) for r in model.requests[:2]]
    assert a["nodes"] == b["nodes"] and a["tables"] == b["tables"]
    assert b["repairFeedback"][0] == BLANK_ROW_ERROR
    detail = json.loads(b["repairFeedback"][1].split(":", 1)[1])["rows"]
    assert [i["row"] for i in detail] == [2] and detail[0]["nonemptySourceRefs"]
    state = states[-1]["tableStages"][a["regionId"]]
    assert state["layout"]["attempts"] == 2
    assert len(state["layout"].get("history", [])) == int(fix)
    assert state["structure"]["attempts"] == int(fix)
    if fix:
        assert result["extraction"]["status"] == "complete"
        assert not any(i.get("tableStage") == "layout" for i in result["issues"])
    run(model, restore=states[-1], maxModelCalls=4)
    assert len(model.requests) == (4 if fix else 2)


def test_cancelled_failed_layout_resumes_with_feedback_not_a_replayed_free_attempt():
    model, states = BlankFirst(), []
    result = run(model, states=states, cancelled=lambda: len(model.requests) == 1)
    assert result["data"] is None and len(model.requests) == 1
    before = states[-1]["tableStages"]["semantic-region:1"]["layout"]
    assert before["attempts"] == 1 and before["feedback"][0] == BLANK_ROW_ERROR
    result = run(model, restore=states[-1])
    assert len(model.requests) == 4 and result["extraction"]["modelCalls"] == 4
    assert (
        json.loads(model.requests[1].messages[-1]["content"])["repairFeedback"]
        == before["feedback"]
    )


@pytest.mark.parametrize("failure", ["id", "key", "contract", "formula"])
def test_mapping_only_error_does_not_reopen_layout_or_reset_any_attempt(failure):
    class BrokenColumns(Model):
        def infer(self, request):
            response = super().infer(request)
            payload = json.loads(request.messages[-1]["content"])
            if payload["tableStage"] == "structure":
                value = json.loads(response.text)
                a, b = value["record"]["columns"].values()
                if failure in {"id", "key"}:
                    b[failure] = a[failure]
                elif failure == "contract":
                    a["invented"] = "not a field"
                else:
                    b.update(bindingMode="formula", valueType="decimal")
                return InferenceResponse(json.dumps(value), response.usage)
            return response

    model, states = BrokenColumns(), []
    result = run(model, states=states)
    assert result["data"] is None
    assert [json.loads(r.messages[-1]["content"])["tableStage"] for r in model.requests] == [
        "layout",
        "structure",
        "structure",
    ]
    state = states[-1]["tableStages"]["semantic-region:1"]
    assert state["layout"]["attempts"] == len(state["layout"]["history"]) == 1
    assert state["structure"]["attempts"] == 2
    assert not state["structure"].get("layoutRevisionPending")
    run(model, restore=states[-1])
    assert len(model.requests) == 3


def test_unlocated_or_forged_diagnostic_cannot_trigger_layout_review():
    doc, region, _ = fixture()
    for error in [
        ValueError("private value"),
        CompileError("duplicate_component_id"),
        CompileError("binding_cannot_represent_requested_type", selection={"bindingId": "foreign"}),
        CompileError("invalid_table_value_selection:fake"),
        CompileError(
            BLANK_ROW_ERROR, selection={"rows": [{"row": 9999, "nonemptySourceRefs": ["fake"]}]}
        ),
    ]:
        assert not needs_layout_review(error, doc, region)


def test_hwpx_bounded_cell_chunks_preserve_spaces_and_add_no_newline_inside_the_original(
    monkeypatch,
):
    from document_files import extractors

    monkeypatch.setattr(extractors, "MAX_UNIT_CHARS", 4)
    original = " \t A  B\nC\t  "
    doc = observe(native_file("hwpx", original), "hwpx")
    table = next(iter(doc.tables.values()))
    cell = next(c for c in table["cells"] if c["row"] == 1)
    composite = doc.nodes[cell["sourceRef"]]
    assert composite["text"] == original
    assert composite["normalization"] == "join_native_cell_segments_preserving_chunks"
    for segment in composite["sourceSegments"]:
        node = doc.nodes[segment["sourceRef"]]
        assert 0 < len(node["text"]) <= 4
        assert composite["text"][segment["start"] : segment["end"]] == node["text"]
    region = prepare_regions(doc, context_chars=16000)[0]
    with pytest.raises(CompileError, match=BLANK_ROW_ERROR):
        layout.accept(response(doc, region, ["header", "blank"]), doc, region)


def test_chunk_separators_require_exact_source_identity_and_adjacent_numbers():
    from document_files.document_model.native import _cell_segment_separator

    a = {"sourceStructure": {"cell": "cell1", "paragraph": 2, "chunk": 1}}
    b = {"sourceStructure": {"cell": "cell1", "paragraph": 2, "chunk": 2}}
    assert _cell_segment_separator(a, b) == ""
    for key, value in [("cell", "cell2"), ("paragraph", 3), ("chunk", 3), ("chunk", True)]:
        changed = copy.deepcopy(b)
        changed["sourceStructure"][key] = value
        assert _cell_segment_separator(a, changed) == "\n"
    assert _cell_segment_separator({"sourceStructure": {}}, {"sourceStructure": {}}) == "\n"


def test_unidentified_chunks_do_not_claim_to_be_one_original_unit():
    from document_files.document_model.native import _cell_segment_separator

    assert _cell_segment_separator(
        {"sourceStructure": {"chunk": 1}}, {"sourceStructure": {"chunk": 2}}
    ) == "\n"


def test_preserved_hwpx_cell_chunks_cannot_overrun_the_native_unit_limit(monkeypatch):
    from document_files import extractors
    from document_files.extraction_errors import ExtractionError
    from document_files.hwpx_structure import _Reader

    monkeypatch.setattr(extractors, "MAX_UNIT_CHARS", 4)
    reader = object.__new__(_Reader)
    reader.units = [None] * 199999
    reader.characters, reader.base = 0, {}
    with pytest.raises(ExtractionError, match="output budget"):
        reader.emit("table_cell", {"cell": "cell", "segment": 1}, "         ")
    assert len(reader.units) == 200000
    assert reader.units[-1].content == "    "
