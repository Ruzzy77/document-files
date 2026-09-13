"""Native columnSpan must survive the observation projection unchanged."""

import io
import zipfile
from copy import deepcopy

import pytest
from openpyxl import Workbook
from test_regional_interpretation import ReferenceModel, run

from document_files.api import AnalysisInput, AnalysisJob, extract_structure_from_stream
from document_files.document_model.native import NATIVE_OBSERVATION_VERSION
from document_files.document_model.observe import observe_document
from document_files.interpretation.engine import observation_identity


def observe(raw, format_id):
    result = extract_structure_from_stream(
        AnalysisJob(
            job_id="native-geometry", input=AnalysisInput.from_bytes(raw, format_id=format_id)
        ),
        io.BytesIO(raw),
    )
    assert result["ok"] and result["completeness"] == "complete"
    nodes = {f"n{n['ordinal']}": n for n in result["units"]}
    before = deepcopy(nodes)
    doc = observe_document(raw, format_id, nodes)
    assert nodes == before and all(doc.nodes[k] == n for k, n in nodes.items())
    assert not doc.issues
    assert doc.provenance["nativeAdapterVersion"] == NATIVE_OBSERVATION_VERSION
    return doc


def assert_geometry(doc):
    table = next(iter(doc.tables.values()))
    assert (table["rowCount"], table["colCount"]) == (3, 3)
    cells = {(c["row"], c["col"]): c for c in table["cells"]}
    assert cells[(0, 0)]["rowSpan"] == 2
    assert cells[(0, 1)]["colSpan"] == cells[(0, 1)]["sourceCell"]["columnSpan"] == 2
    assert set(cells) == {(0, 0), (0, 1), (1, 1), (1, 2), (2, 0), (2, 1), (2, 2)}
    assert len(cells) == 7  # No duplicate anchor or fabricated covered cell.


def test_actual_xlsx_merged_rows_and_columns_keep_original_geometry():
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "Measurements"])
    sheet.append([None, "Length", "Width"])
    sheet.append(["Q-0007", "001.2300", None])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    sheet["C3"].number_format = "@"  # An actually stored blank, not a missing cell.
    out = io.BytesIO()
    book.save(out)
    book.close()
    doc = observe(out.getvalue(), "xlsx")
    assert_geometry(doc)
    table = next(iter(doc.tables.values()))
    by_coordinate = {c["sourceCell"]["coordinate"]: c for c in table["cells"]}
    for coordinate, expected in [("A3", "Q-0007"), ("B3", "001.2300")]:
        node = doc.nodes[by_coordinate[coordinate]["sourceRef"]]
        assert node["semantic"]["value"]["value"] == expected
    assert doc.nodes[by_coordinate["C3"]["sourceRef"]]["semantic"]["value"]["kind"] == "blank"


def test_actual_hwpx_merged_header_geometry_from_original_cellspan_xml():
    rows = [
        [(0, 2, 1, "ID"), (1, 1, 2, "Measurements")],
        [(1, 1, 1, "Length"), (2, 1, 1, "Width")],
        [(0, 1, 1, "Q-0007"), (1, 1, 1, "001.2300"), (2, 1, 1, "")],
    ]
    body = '<sec><p><run><tbl rowCnt="3" colCnt="3">'
    for row, cells in enumerate(rows):
        body += "<tr>"
        for col, height, width, text in cells:
            body += (
                f'<tc header="{int(row < 2)}"><cellAddr rowAddr="{row}" colAddr="{col}"/>'
                f'<cellSpan rowSpan="{height}" colSpan="{width}"/>'
                f"<subList><p><run><t>{text}</t></run></p></subList></tc>"
            )
        body += "</tr>"
    body += "</tbl></run></p></sec>"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("Contents/section0.xml", body)
    doc = observe(out.getvalue(), "hwpx")
    assert_geometry(doc)
    table = next(iter(doc.tables.values()))
    data = sorted((c for c in table["cells"] if c["row"] == 2), key=lambda c: c["col"])
    assert [doc.nodes[c["sourceRef"]]["text"] for c in data] == ["Q-0007", "001.2300", ""]


@pytest.mark.parametrize(
    "geometry,expected", [({}, 1), ({"colSpan": 2}, 2), ({"columnSpan": 3, "colSpan": 2}, 3)]
)
def test_native_public_name_takes_precedence_without_changing_legacy_nodes(geometry, expected):
    node = {
        "text": "Value",
        "semantic": {
            "sheet": {"index": 1, "name": "One"},
            "cell": {"row": 1, "column": 1, "indexBase": 1, **geometry},
        },
    }
    original = deepcopy(node)
    doc = observe_document(b"", "xlsx", {"n": node})
    assert next(iter(doc.tables.values()))["cells"][0]["colSpan"] == expected
    assert doc.nodes["n"] == node == original


def test_native_adapter_identity_rejects_old_geometry_checkpoint_without_model_call():
    assert observation_identity(None)["version"] == NATIVE_OBSERVATION_VERSION
    saved = []
    run(model=ReferenceModel(), checkpoint=saved.append)
    old = deepcopy(saved[-1])
    old["identity"]["observationBackend"]["version"] = "document-files.observation.v1"
    model = ReferenceModel()
    with pytest.raises(ValueError, match="incompatible"):
        run(model=model, restore=old)
    assert model.calls == 0
