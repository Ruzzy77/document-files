"""Lossless table requests and actual-context replanning, not AI quality claims."""

import io
import json
import zipfile
from copy import deepcopy

import pytest
from openpyxl import Workbook
from test_native_merged_geometry import observe
from test_table_protocol import TableModel, execute, record_response

from document_files.document_model.observe import observe_document
from document_files.interpretation.backends import InferenceResponse
from document_files.interpretation.compiler import compile_region
from document_files.interpretation.legacy_engine import contract_messages
from document_files.interpretation.regions import (
    add_table_definition_context,
    column_candidates,
    prepare_regions,
    region_payload,
    replan_table_region,
)
from document_files.interpretation.source_dictionary import _same
from document_files.interpretation.table_protocol import (
    STRUCTURE_SYSTEM,
    meaning_decision_schema,
    meaning_payload,
    structural_ir,
    structure_payload,
    structure_schema,
)
from document_files.interpretation.table_reference_wire import prepare_meaning_wire
from document_files.interpretation.table_selection import SYSTEM as SELECTION_SYSTEM
from document_files.interpretation.table_selection_wire import selection_schema
from document_files.interpretation.table_source_wire import (
    SYSTEM,
    compact_table_sources,
    expand_table_sources,
)
from document_files.interpretation.table_sources import source_inventory


def encoded(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


@pytest.mark.parametrize("special", [None, {}, [], True, 1, 1.0, "", " \t", "001.2300"])
def test_every_source_type_missing_key_duplicate_record_and_path_is_preserved(special):
    payload = {"nodes": {}, "tables": {}, "relations": []}
    for i in range(40):
        text = "동일한 값" if i in (11, 12) else f"원문 {i}"
        payload["nodes"][f"n{i}"] = {
            "text": text,
            "semantic": {"value": {"raw": text, "value": special}, "duplicate": special},
            "sourceStructure": {"shared": "원래 서식 참조" * 30, "a/b~c": {"": i}},
            **({"optional": special} if i % 2 else {}),
        }
        payload["relations"].append({"kind": "contains", "sourceRef": "root", "targetRef": f"n{i}"})
    before = deepcopy(payload)
    packed = compact_table_sources(payload)
    assert "tableSourceEncoding" in packed
    assert len(encoded(packed)) + len(SYSTEM) < len(encoded(payload))
    assert list(packed["nodes"]) == list(payload["nodes"])
    assert [n["text"] for n in packed["nodes"].values()] == [
        n["text"] for n in payload["nodes"].values()
    ]
    assert _same(expand_table_sources(packed), payload)
    assert _same(payload, before)
    assert SYSTEM in contract_messages(STRUCTURE_SYSTEM, packed, {})[0]["content"]


def test_equal_python_numbers_do_not_share_typed_columns_or_rows():
    payload = {
        "nodes": {
            f"n{i}": {
                "text": "",
                "sourceStructure": {
                    "shared": "long enough metadata " * 20,
                    "bool": True,
                    "int": 1,
                    "float": 1.0,
                    "vary": [True, 1, 1.0][i % 3],
                },
            }
            for i in range(30)
        }
    }
    assert _same(expand_table_sources(compact_table_sources(payload)), payload)


def test_small_reserved_and_already_packed_requests_are_not_silently_reinterpreted():
    for payload in [
        {"nodes": {}},
        {"nodes": {"a": {"text": "same"}, "b": {"text": "same"}}},
        {"nodes": {"a": {"metadata": ["source property"]}}},
    ]:
        assert compact_table_sources(payload) is payload
    with pytest.raises(ValueError, match="already_encoded"):
        compact_table_sources({"tableSourceEncoding": "existing"})
    with pytest.raises(ValueError, match="unknown"):
        expand_table_sources({"tableSourceEncoding": "future"})


def native_table(format_id, count=50, leading=False):
    rows = [
        ["ID", "Length", "Width"],
        *[
            [
                f"Q-{i:04}",
                "001.2300" if i in (11, 12) else f"{i}.0100",
                "" if i == 37 else "002.3400",
            ]
            for i in range(1, count + 1)
        ],
    ]
    out = io.BytesIO()
    if format_id == "xlsx":
        book = Workbook()
        sheet = book.active
        if leading:
            for index, text in enumerate(["검체 치수", "단위: mm", "검체별 기록"], 1):
                sheet.append([text])
                sheet.merge_cells(start_row=index, end_row=index, start_column=1, end_column=3)
        for row in rows:
            sheet.append(row)
        book.save(out)
        book.close()
    else:
        body = f'<sec><p><run><tbl rowCnt="{len(rows)}" colCnt="3">'
        for row, values in enumerate(rows):
            body += "<tr>"
            for col, value in enumerate(values):
                body += (
                    f'<tc><cellAddr rowAddr="{row}" colAddr="{col}"/>'
                    '<cellSpan rowSpan="1" colSpan="1"/>'
                    f"<subList><p><run><t>{value}</t></run></p></subList></tc>"
                )
            body += "</tr>"
        body += "</tbl></run></p></sec>"
        with zipfile.ZipFile(out, "w") as package:
            package.writestr("Contents/section0.xml", body)
    return out.getvalue()


@pytest.mark.parametrize("format_id", ["hwpx", "xlsx"])
def test_actual_native_table_requests_keep_every_cell_source_and_metadata(format_id):
    doc = observe(native_table(format_id), format_id)
    before = deepcopy(doc)
    regions = prepare_regions(doc, context_chars=16000)
    assert len(regions) < 12  # Reproduced native HWPX previously used about one row per request.
    seen = []
    for region in regions:
        original = region_payload(doc, region)
        packed = structure_payload(original)
        restored = expand_table_sources(packed)
        assert _same(original["nodes"], restored["nodes"])
        assert _same(original["relations"], restored["relations"])
        for ref, table in original["tables"].items():
            assert _same(table["cells"], restored["tables"][ref]["cells"])
            seen.extend(cell["sourceRef"] for cell in doc.tables[ref]["cells"])
        assert [n["text"] for n in packed["nodes"].values()] == [
            n["text"] for n in original["nodes"].values()
        ]
        messages = contract_messages(STRUCTURE_SYSTEM, packed, structure_schema(doc, region))
        # prepare_regions uses the same default metadata when sizing its request.
        planned = structure_payload(original | {"intent": "discover", "targetHandles": {}})
        measured = sum(
            len(m["content"])
            for m in contract_messages(STRUCTURE_SYSTEM, planned, structure_schema(doc, region, {}))
        )
        assert region["requestChars"] == measured <= 16000
        assert sum(len(m["content"]) for m in messages) <= 16000
    original_refs = [c["sourceRef"] for table in before.tables.values() for c in table["cells"]]
    assert seen == original_refs and len(seen) == len(set(seen))
    assert doc.nodes == before.nodes and doc.bindings == before.bindings


def long_html():
    return (
        "<table><tr><th>Code</th><th>Size</th></tr>"
        + "".join(f"<tr><td>{i:04}</td><td>1.2300</td></tr>" for i in range(50))
        + "</table>"
    ).encode()


def test_dynamic_replanning_preserves_original_table_headers_rows_and_mapping():
    doc = observe_document(long_html(), "html", {})
    regions = prepare_regions(doc, context_chars=11000)
    region = regions[1]
    original = deepcopy(doc.tables[region["tableRef"]])
    metadata = {
        "intent": "discover",
        "targetHandles": {},
        "sameTableMapping": {"label": "same source label " * 150},
    }
    replacements = replan_table_region(doc, region, context_chars=11000, request_metadata=metadata)
    assert replacements and len(replacements) > 1
    cells = []
    for child in replacements:
        assert child["id"].startswith(region["id"] + ":part:")
        table = doc.tables[child["tableRef"]]
        assert table["sourceTableRef"] == original["sourceTableRef"]
        assert table["headerCells"] == original["headerCells"]
        cells.extend(table["cells"])
        request = structure_payload(region_payload(doc, child) | metadata)
        assert expand_table_sources(request)["sameTableMapping"] == metadata["sameTableMapping"]
        size = sum(
            len(m["content"])
            for m in contract_messages(STRUCTURE_SYSTEM, request, structure_schema(doc, child, {}))
        )
        assert size == child["requestChars"] <= 11000
    assert cells == original["cells"]


def test_prior_mapping_keeps_real_xlsx_headers_after_title_without_declaring_them():
    doc = observe(native_table("xlsx", leading=True), "xlsx")
    regions = prepare_regions(doc, context_chars=16000)
    region = regions[1]
    root_ref = doc.tables[region["tableRef"]]["sourceTableRef"]
    root = deepcopy(doc.tables[root_ref])
    header_cells = [c for c in root["cells"] if c["row"] == 3]
    refs = [c["sourceRef"] for c in header_cells]
    assert not set(refs).intersection(region["contextNodeIds"])
    before_nodes = deepcopy(doc.nodes)
    owned = list(region["nodeIds"])
    bindings = list(region["bindingIds"])
    required = list(region["requiredBindingIds"])
    add_table_definition_context(doc, region, refs)
    assert set(refs) <= set(region["contextNodeIds"])
    assert region["nodeIds"] == owned and region["bindingIds"] == bindings
    assert region["requiredBindingIds"] == required
    view = doc.tables[region["tableRef"]]
    assert set(refs) <= {c["sourceRef"] for c in view["leadingCells"]}
    assert all(not c.get("isHeader") for c in view["leadingCells"])
    assert all(not c["headerRefs"] for c in column_candidates(view)[0])
    assert doc.tables[root_ref] == root and doc.nodes == before_nodes
    packed = structure_payload(region_payload(doc, region))
    restored = expand_table_sources(packed)
    assert all(restored["nodes"][ref]["text"] == doc.nodes[ref]["text"] for ref in refs)
    state = deepcopy(doc), deepcopy(region)
    add_table_definition_context(doc, region, refs)
    assert (doc, region) == state


class LongMappingModel(TableModel):
    def infer(self, request):
        payload = json.loads(request.messages[-1]["content"])
        if payload.get("tableStage") == "structure":
            self.requests.append(request)
            restored = expand_table_sources(payload)
            table = next(iter(restored["tables"].values()))
            cells = table["cells"]
            if isinstance(cells, dict):
                cells = [dict(zip(cells["columns"], row, strict=True)) for row in cells["rows"]]
            headers = table.get("headerCells", [])
            if isinstance(headers, dict):
                headers = [
                    dict(zip(headers["columns"], row, strict=True)) for row in headers["rows"]
                ]
            table["cells"] = [*cells, *headers]
            response = record_response(restored)
            response["record"]["rowStart"] = min(c["row"] for c in cells)
            for column in response["record"]["columns"]:
                column["label"] = column["label"] * 90
            return InferenceResponse(
                json.dumps(response), {"prompt_tokens": 10, "completion_tokens": 20}
            )
        return super().infer(request)


def test_engine_replans_before_call_and_checkpoint_resume_never_replays_accepted_rows():
    model, states = LongMappingModel(), []
    content = long_html()
    result = execute(model, content=content, states=states, contextChars=11000, maxModelCalls=60)
    assert not any(i["code"] == "region_context_budget_exceeded" for i in result["issues"]), result[
        "issues"
    ]
    assert result["data"]["records"] == [{"code": f"{i:04}", "size": "1.2300"} for i in range(50)]
    assert any(":part:" in region["id"] for region in states[-1]["regions"])
    calls = len(model.requests)
    resumed = execute(
        model, content=content, restore=states[-1], contextChars=11000, maxModelCalls=60
    )
    assert len(model.requests) == calls and resumed["data"] == result["data"]
    assert all(sum(len(m["content"]) for m in req.messages) <= 11000 for req in model.requests)


def test_partial_checkpoint_keeps_new_slices_without_replaying_the_first_structure():
    model, states = LongMappingModel(), []
    content = long_html()
    result = execute(model, content=content, states=states, contextChars=11000, maxModelCalls=2)
    assert result["extraction"]["status"] == "partial"
    assert len(model.requests) == 2 and len(states[-1]["accepted"]) == 1
    assert any(":part:" in region["id"] for region in states[-1]["regions"])
    prior_region = next(iter(states[-1]["accepted"]))
    resumed = execute(
        model,
        content=content,
        restore=states[-1],
        contextChars=11000,
        maxModelCalls=2,
        additional_budget={"maxModelCalls": 58},
    )
    assert resumed["data"]["records"] == [{"code": f"{i:04}", "size": "1.2300"} for i in range(50)]
    assert all(
        json.loads(req.messages[-1]["content"])["regionId"] != prior_region
        for req in model.requests[2:]
    )


def test_an_indivisible_table_row_does_not_create_an_endless_replan_or_extra_view():
    doc = observe_document(b"<table><tr><td>one row</td></tr></table>", "html", {})
    region = prepare_regions(doc, context_chars=16000)[0]
    before = deepcopy(doc)
    metadata = {
        "intent": "discover",
        "targetHandles": {},
        "sameTableMapping": {"label": "x" * 17000},
    }
    assert replan_table_region(doc, region, context_chars=16000, request_metadata=metadata) is None
    assert doc == before


@pytest.mark.parametrize(
    "key,old",
    [
        ("tableProtocolVersion", "document-files.table-protocol.v20"),
        ("tableProtocolVersion", "document-files.table-protocol.v21"),
        ("regionPlanVersion", "document-files.region-plan.v20"),
    ],
)
def test_old_table_display_and_planning_checkpoints_are_rejected_before_call(key, old):
    model, states = TableModel(), []
    execute(model, states=states)
    state = deepcopy(states[-1])
    state["identity"][key] = old
    calls = len(model.requests)
    with pytest.raises(ValueError, match="incompatible"):
        execute(model, restore=state)
    assert len(model.requests) == calls


@pytest.mark.parametrize("format_id", ["hwpx", "xlsx"])
def test_meaning_display_preserves_every_source_choice_and_frozen_native_record(format_id):
    doc = observe(native_table(format_id), format_id)
    region = prepare_regions(doc, context_chars=16000)[0]
    table = doc.tables[region["tableRef"]]
    headers = {c["col"]: c["sourceRef"] for c in table["cells"] if c["row"] == 0}
    response = {
        "regionId": region["id"],
        "tableKind": "record_table",
        "record": {
            "id": "rows",
            "key": "rows",
            "label": "Rows",
            "tableRef": region["tableRef"],
            "rowStart": 0,
            "rowEnd": max(c["row"] for c in table["cells"]),
            "definitionRefs": list(headers.values()),
            "rowRoles": [
                {"row": row, "role": "header" if row == 0 else "data"}
                for row in sorted({c["row"] for c in table["cells"]})
            ],
            "columns": [
                {
                    "id": key,
                    "key": key,
                    "label": key,
                    "column": col,
                    "valueType": "string",
                    "definitionRefs": [headers[col]],
                }
                for col, key in enumerate(["id", "length", "width"])
            ],
        },
    }
    _, ir = structural_ir(response, doc, region)
    compiled = compile_region(ir, doc, region)
    payload = meaning_payload(
        region_payload(doc, region), ir, compiled, source_inventory(doc, region)
    )
    wire = prepare_meaning_wire(payload, meaning_decision_schema(doc, region, ir, {}))
    original = wire.payload | {"meaningPhase": "selection"}
    schema = selection_schema(original["meaningSources"])
    packed = compact_table_sources(original)
    assert _same(expand_table_sources(packed), original)
    assert "nodes" not in packed  # No invented source dictionary in a meaning request.
    assert packed["meaningSources"] == original["meaningSources"]
    assert packed["frozenStructure"] == original["frozenStructure"]
    assert selection_schema(packed["meaningSources"]) == schema
    before = sum(len(m["content"]) for m in contract_messages(SELECTION_SYSTEM, original, schema))
    after = sum(len(m["content"]) for m in contract_messages(SELECTION_SYSTEM, packed, schema))
    assert after < before
    # Sharing cannot guarantee every later meaning/repair fits; never drop sources
    # or broaden the configured limit to make that separate assertion true.


def test_resume_uses_only_preceding_mappings_not_later_accepted_regions(monkeypatch):
    from document_files.interpretation import engine

    model, states = LongMappingModel(invalid_meaning=True), []
    content = long_html()
    execute(model, content=content, states=states, contextChars=11000, maxModelCalls=60)
    assert len(states[-1]["accepted"]) > 1
    first = states[-1]["regions"][0]["id"]
    assert states[-1]["tableStages"][first]["meaning"]["status"] != "complete"
    seen = []
    original = engine.add_table_definition_context

    def record(doc, region, refs):
        seen.append(region["id"])
        return original(doc, region, refs)

    monkeypatch.setattr(engine, "add_table_definition_context", record)
    model.invalid_meaning = False
    resumed = execute(
        model,
        content=content,
        restore=states[-1],
        contextChars=11000,
        maxModelCalls=60,
        additional_budget={"maxModelCalls": 20},
    )
    assert seen and first not in seen
    assert resumed["data"]["records"] == [{"code": f"{i:04}", "size": "1.2300"} for i in range(50)]


def test_later_slice_receives_the_nearest_accepted_mapping_not_the_first_one():
    class ChangingLabelModel(LongMappingModel):
        last_label = None
        structure_calls = 0

        def infer(self, request):
            payload = json.loads(request.messages[-1]["content"])
            response = super().infer(request)
            if payload.get("tableStage") != "structure":
                return response
            if self.last_label is not None:
                assert payload["sameTableMapping"]["label"] == self.last_label
            self.structure_calls += 1
            self.last_label = "Scripted region label " + payload["regionId"]
            value = json.loads(response.text)
            value["record"]["label"] = self.last_label
            return InferenceResponse(json.dumps(value), response.usage)

    model = ChangingLabelModel()
    result = execute(model, content=long_html(), contextChars=11000, maxModelCalls=60)
    assert model.structure_calls >= 3
    assert result["data"]["records"] == [{"code": f"{i:04}", "size": "1.2300"} for i in range(50)]
