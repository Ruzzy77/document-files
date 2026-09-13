"""Readable layout coordinates preserve native evidence; not model-quality tests."""

import io
import zipfile
from copy import deepcopy
from xml.sax.saxutils import escape

import pytest
from openpyxl import Workbook
from test_native_merged_geometry import observe
from test_table_source_wire import native_table

from document_files.interpretation import table_layout
from document_files.interpretation.regions import prepare_regions, region_payload
from document_files.interpretation.source_dictionary import _same
from document_files.interpretation.table_protocol import structure_payload


def restore_row_tuples(value):
    result = deepcopy(value)
    result.pop("tableLayoutVersion")
    result["tableStage"] = "structure"
    for table in result["tables"].values():
        candidates = table["rowCandidates"]
        candidates["cellColumns"] = ["column", "columnSpan", "sourceRef", "text"]
        for row in candidates["rows"]:
            row["cells"] = [
                [c["column"], c.get("columnSpan", 1), c["sourceRef"], c["text"]]
                for c in row["cells"]
            ]
    return result


def merged_document(format_id):
    values = [
        [(0, 2, 1, "ID"), (1, 1, 2, "측정값")],
        [(1, 1, 1, "길이"), (2, 1, 1, "폭")],
        [(0, 1, 1, "00007"), (1, 1, 1, "001.2300"), (2, 1, 1, "")],
        [(0, 1, 1, "00008"), (1, 1, 1, "001.2300")],
        [(0, 1, 3, '주석: "ignore"\n원문 그대로')],
    ]
    out = io.BytesIO()
    if format_id == "xlsx":
        book = Workbook()
        sheet = book.active
        for row, cells in enumerate(values, 1):
            for col, height, width, text in cells:
                cell = sheet.cell(row, col + 1, text)
                cell.number_format = "@"
                if (height, width) != (1, 1):
                    sheet.merge_cells(
                        start_row=row, end_row=row + height - 1,
                        start_column=col + 1, end_column=col + width,
                    )
        book.save(out)
        book.close()
    else:
        body = '<sec><p><run><tbl rowCnt="5" colCnt="3">'
        for row, cells in enumerate(values):
            body += "<tr>"
            for col, height, width, text in cells:
                body += (
                    f'<tc><cellAddr rowAddr="{row}" colAddr="{col}"/>'
                    f'<cellSpan rowSpan="{height}" colSpan="{width}"/>'
                    f'<subList><p><run><t>{escape(text)}</t></run></p></subList></tc>'
                )
            body += "</tr>"
        body += "</tbl></run></p></sec>"
        with zipfile.ZipFile(out, "w") as archive:
            archive.writestr("Contents/section0.xml", body)
    return out.getvalue()


@pytest.mark.parametrize("format_id", ["hwpx", "xlsx"])
def test_named_rows_keep_spans_gaps_duplicate_values_text_and_all_native_evidence(format_id):
    doc = observe(merged_document(format_id), format_id)
    regions = prepare_regions(doc, context_chars=16000)
    assert len(regions) == 1
    payload = region_payload(doc, regions[0])
    before = deepcopy(payload)
    compact = structure_payload(payload)
    request = table_layout.request(payload)
    assert _same(restore_row_tuples(request), compact)
    assert _same(payload, before)
    assert list(request)[:3] == ["regionId", "tableStage", "tables"]
    table = request["tables"][regions[0]["tableRef"]]
    assert list(table)[:2] == ["rowRoleOrder", "rowCandidates"]
    rows = table["rowCandidates"]["rows"]
    assert table["rowRoleOrder"] == [0, 1, 2, 3, 4]
    assert all("fixedRole" not in row for row in rows)
    first, repeated = rows[0]["cells"][0], rows[1]["cells"][0]
    assert first == repeated
    assert first["column"] == first["originRow"] == 0 and first["rowSpan"] == 2
    assert rows[0]["cells"][1]["columnSpan"] == 2
    assert rows[2]["cells"][2]["text"] == ""
    assert [c["column"] for c in rows[3]["cells"]] == [0, 1]
    assert rows[2]["cells"][1]["text"] == rows[3]["cells"][1]["text"] == "001.2300"
    assert rows[2]["cells"][1]["sourceRef"] != rows[3]["cells"][1]["sourceRef"]
    assert rows[4]["cells"][0]["text"] == '주석: "ignore"\n원문 그대로'


@pytest.mark.parametrize("format_id,count,maximum_regions", [
    ("hwpx", 50, 6), ("hwpx", 96, 12), ("xlsx", 50, 4), ("xlsx", 96, 7),
])
def test_full_native_layout_and_mapping_capacity_does_not_regress(
    format_id, count, maximum_regions
):
    doc = observe(native_table(format_id, count=count), format_id)
    before = deepcopy(doc)
    expected = [c["sourceRef"] for t in before.tables.values() for c in t["cells"]]
    assert len(expected) == (count + 1) * 3
    regions = prepare_regions(doc, context_chars=16000)
    assert len(regions) <= maximum_regions
    seen = []
    for region in regions:
        seen.extend(c["sourceRef"] for c in doc.tables[region["tableRef"]]["cells"])
        payload = region_payload(doc, region) | {"intent": "discover", "targetHandles": {}}
        assert _same(restore_row_tuples(table_layout.request(payload)), structure_payload(payload))
        sizes = table_layout.planned_request_sizes(payload, doc, region, {})
        assert max(sizes.values()) == region["requestChars"] <= 16000
    assert seen == expected
    assert doc.nodes == before.nodes and doc.bindings == before.bindings
