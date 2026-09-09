"""Structural contract tests, not an AI or OCR quality qualification."""

from copy import deepcopy

from document_files.document_model.observe import observe_document


def resolve(doc, binding):
    source = doc.nodes[binding["sourceRef"]]
    for token in binding["path"].split("/")[1:]:
        source = source[token]
    return source[binding["start"] : binding["end"]] if binding["start"] is not None else source


def test_legacy_nodes_preserved_and_exact_zero_has_distinct_location():
    content = b"Reference: 00091\r\nCount: 0\r\nDeadline:\r\n\r\n"
    legacy = {"n1": {"text": content.decode().rstrip(), "sourceStructure": {"paragraph": 1}}}
    before = deepcopy(legacy)
    doc = observe_document(content, "txt", legacy)
    assert legacy == before
    assert doc.nodes["n1"] == before["n1"]
    values = [b for b in doc.bindings.values() if b.get("candidateRole") == "value"]
    zero = next(
        b
        for b in values
        if resolve(doc, b) == "0" and b["sourceRef"].startswith("source:text/line/")
    )
    assert doc.nodes[zero["sourceRef"]]["text"] == "Count: 0"
    assert resolve(doc, doc.bindings[zero["labelRefs"][0]]) == "Count"
    blank = next(
        b for b in values if b.get("blank") and b["sourceRef"].startswith("source:text/line/")
    )
    assert blank["start"] == blank["end"]
    assert doc.nodes[blank["sourceRef"]]["text"] == "Deadline:"
    assert any(n.get("semanticRole") == "blank" for n in doc.nodes.values())
    assert doc.nodes["source:text/line/1"]["sourceStructure"]["lineEnding"] == "\r\n"
    assert doc.to_dict()["schemaVersion"] == "document-files.observation.v1"


def test_markdown_tables_and_footnotes_are_syntax_not_typography():
    content = b"# Stock\n\n| Item | Count |\n| --- | --- |\n| Seed[^q] | 8 |\n\n[^q]: Keep dry.\n"
    doc = observe_document(content, "md", {})
    table = next(t for t in doc.tables.values() if t["basis"] == "markdown_table_tokens")
    assert table["rowCount"] == 2 and table["colCount"] == 2
    assert len(table["cells"]) == 4
    assert sum(c["isHeader"] for c in table["cells"]) == 2
    links = [r for r in doc.relations if r["kind"] == "noteReference"]
    assert any(doc.nodes[r["targetRef"]]["text"] == "Keep dry." for r in links)
    assert doc.provenance["markdown"]["typographyInference"] is False


def test_html_merged_headers_empty_cells_and_form_labels():
    content = b"""<h1>Inventory</h1><table><tr><th rowspan="2">Lot</th><th colspan="2">mL</th></tr>
<tr><th>Start</th><th>Left</th></tr><tr><td>X</td><td>7.50</td><td></td></tr></table>
<label for="ok">Checked</label><input id="ok" type="checkbox" checked>
<label for="date">Date</label><input id="date" value="">"""
    doc = observe_document(content, "html", {})
    table = next(iter(doc.tables.values()))
    assert len(table["cells"]) == 7  # No synthetic duplicate cells for merged header positions.
    assert table["rowCount"] == 3 and table["colCount"] == 3
    assert table["cells"][0]["rowSpan"] == 2
    assert table["cells"][1]["colSpan"] == 2
    assert table["cells"][-1]["row"] == 2 and table["cells"][-1]["col"] == 2
    assert doc.nodes[table["cells"][-1]["sourceRef"]]["text"] == ""
    assert len([r for r in doc.relations if r["kind"] == "labelFor"]) == 2
    assert any(resolve(doc, b) is True for b in doc.bindings.values() if b["path"] != "/text")


def test_native_sheet_coordinates_and_precise_bindings_are_additive():
    legacy = {
        "n1": {
            "text": "A2=0.1",
            "sourceStructure": {},
            "semantic": {
                "sheet": {"index": 1, "name": "Ledger"},
                "cell": {"row": 2, "column": 1, "indexBase": 1},
                "value": {"kind": "number", "value": 0.1, "raw": "0.10000000000000000001"},
            },
        }
    }
    doc = observe_document(b"native bytes", "xlsx", legacy)
    table = next(iter(doc.tables.values()))
    assert table["cells"][0]["row"] == 1 and table["cells"][0]["col"] == 0
    assert any(
        resolve(doc, b) == "0.10000000000000000001"
        for b in doc.bindings.values()
        if b["path"] == "/semantic/value/raw"
    )
    assert doc.nodes["n1"] == legacy["n1"]
    assert doc.regions[0]["tableRef"] == table["id"]


def test_native_multiparagraph_cell_is_explicit_composite_not_first_segment(tmp_path):
    from docx import Document

    from document_files.engine import extract_structure

    source = tmp_path / "paragraphs.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "First clause"
    cell.add_paragraph("Second clause")
    document.save(source)
    native = extract_structure(source)
    nodes = {f"n{n['ordinal']}": n for n in native["units"]}
    before = deepcopy(nodes)
    observed = observe_document(source.read_bytes(), "docx", nodes)
    table = next(iter(observed.tables.values()))
    assert len(table["cells"]) == 1
    cell = table["cells"][0]
    composite = observed.nodes[cell["sourceRef"]]
    assert composite["text"] == "First clause\nSecond clause"
    assert composite["observationBasis"] == "program_normalized"
    assert len(composite["sourceSegments"]) == 2
    for segment in composite["sourceSegments"]:
        assert (
            composite["text"][segment["start"] : segment["end"]]
            == nodes[segment["sourceRef"]]["text"]
        )
    assert nodes == before
    assert all(observed.nodes[key] == original for key, original in before.items())


def test_native_note_type_prevents_cross_linking_equal_comment_and_footnote_ids():
    nodes = {
        "n1": {"text": "", "sourceStructure": {"reference_type": "footnoteReference", "note": "1"}},
        "n2": {"text": "Footnote", "sourceStructure": {"container_kind": "footnote", "note": "1"}},
        "n3": {"text": "Comment", "sourceStructure": {"container_kind": "comment", "note": "1"}},
    }
    observed = observe_document(b"native document", "docx", nodes)
    links = [r for r in observed.relations if r["kind"] == "noteReference"]
    assert links == [
        {"kind": "noteReference", "sourceRef": "n1", "targetRef": "n2", "basis": "native_reference"}
    ]


def test_nested_native_table_has_separate_cells_and_container_relation(tmp_path):
    from docx import Document

    from document_files.engine import extract_structure

    source = tmp_path / "nested.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1).cell(0, 0)
    outer.text = "Outer note"
    outer.add_table(rows=1, cols=1).cell(0, 0).text = "Inner value"
    document.save(source)
    result = extract_structure(source)
    nodes = {f"n{n['ordinal']}": n for n in result["units"]}
    observed = observe_document(source.read_bytes(), "docx", nodes)
    assert len(observed.tables) == 2
    texts = [
        observed.nodes[cell["sourceRef"]]["text"]
        for table in observed.tables.values()
        for cell in table["cells"]
    ]
    assert "Outer note" in texts and "Inner value" in texts
    assert all(not ("Outer note" in text and "Inner value" in text) for text in texts)
    assert any(r["kind"] == "nestedTableInCell" for r in observed.relations)


def test_deep_html_reports_partial_instead_of_unbounded_tree():
    content = ("<div>" * 300 + "text" + "</div>" * 300).encode()
    observed = observe_document(content, "html", {})
    assert observed.coverage["status"] == "partial"
    assert any(i["code"] == "observation_budget_exceeded" for i in observed.issues)


def test_line_context_does_not_repeat_full_document_container():
    text = "".join(f"Field {i}: {i}\n" for i in range(500))
    doc = observe_document(text.encode(), "txt", {})
    assert doc.nodes["source:text"]["text"] == text
    assert all("source:text" not in r["contextNodeIds"] for r in doc.regions)
    assert sum(r.get("sourceRef") == "source:text" for r in doc.relations) == 500


def test_bounded_regions_omit_full_source_and_unrelated_contains_edges():
    from document_files.interpretation.regions import prepare_regions, region_payload

    text = "".join(f"Field {i}: {i}\n" for i in range(500))
    doc = observe_document(text.encode(), "txt", {})
    regions = prepare_regions(doc, context_chars=20000)
    assert regions and all(r["withinContextBudget"] for r in regions)
    first = region_payload(doc, regions[0])
    assert "source:text" not in first["nodes"]
    assert all(r.get("sourceRef") != "source:text" for r in first["relations"])
    assert doc.nodes["source:text"]["text"] == text


def test_table_adjacent_context_is_candidate_and_source_order_is_preserved():
    html = (
        b"<h1>Inventory</h1><p>All amounts in kg.</p>"
        b"<table><tr><th>Item</th><th>Amount</th></tr>"
        b"<tr><td>A</td><td>1</td></tr></table><p>Only sealed items.</p>"
    )
    doc = observe_document(html, "html", {})
    table_region = next(r for r in doc.regions if r.get("tableRef"))
    context = [doc.nodes[n]["text"] for n in table_region["contextNodeIds"]]
    assert "Inventory" in context and "All amounts in kg." in context
    assert "Only sealed items." in context
    assert not doc.regions[0].get("tableRef")
    candidates = [r for r in doc.relations if r.get("kind") == "contextCandidate"]
    assert candidates and all(r["applicability"] == "not_assessed" for r in candidates)


def test_row_views_retain_cells_spanning_from_prior_rows():
    from document_files.document_model.model import ObservationDocument
    from document_files.document_model.native import bind_spans
    from document_files.interpretation.regions import prepare_regions

    doc = ObservationDocument()
    shared = doc.node("shared", "Batch A")
    bind_spans(doc, shared)
    cells = [{"sourceRef": shared, "row": 0, "col": 0, "rowSpan": 32, "colSpan": 1}]
    for row in range(32):
        node = doc.node(f"cell:{row}", f"{row} " + "x" * 350)
        bind_spans(doc, node)
        cells.append({"sourceRef": node, "row": row, "col": 1, "rowSpan": 1, "colSpan": 1})
    doc.tables["t"] = {
        "id": "t",
        "basis": "native_structure",
        "cells": cells,
        "rowCount": 32,
        "colCount": 2,
    }
    doc.regions = [
        {
            "id": "r",
            "tableRef": "t",
            "nodeIds": list(doc.nodes),
            "contextNodeIds": [],
            "bindingIds": list(doc.bindings),
        }
    ]
    regions = prepare_regions(doc, context_chars=15000)
    assert len(regions) > 1
    later = [
        doc.tables[r["tableRef"]]
        for r in regions
        if doc.tables[r["tableRef"]].get("viewRowStart", 0) > 0
    ]
    assert later
    assert all(any(c["sourceRef"] == "shared" for c in t["cells"]) for t in later)
    assert all(t["viewRowStart"] <= t["viewRowEnd"] for t in later)
