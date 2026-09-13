"""Stored blank cells are source observations, not implicit grid gaps or nulls."""

import io
import random
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Side
from test_native_merged_geometry import observe

from document_files.interpretation.compiler import _read, preferred_binding
from document_files.xlsx_empty_cells import MergeCoverage, stored_empty_type


@pytest.mark.parametrize(
    "body,expected",
    [
        ("<c/>", "n"),
        ("<c><v/></c>", "n"),
        ('<c t="inlineStr"><is><t/></is></c>', "inlineStr"),
        ('<c t="str"><v/></c>', "str"),
        ("<c><v>0</v></c>", None),
        ("<c><f>1+1</f><v/></c>", None),
        ('<c><f t="shared" si="1"/></c>', None),
        ('<c t="s"/>', None),
        ('<c t="e"/>', None),
        ('<c t="b"/>', None),
        ('<c t="inlineStr"><is><r><t> </t></r></is></c>', None),
    ],
)
def test_only_source_empty_cell_types_are_blank(body, expected):
    root = ET.fromstring(body.replace("<c", '<c xmlns="urn:sheet"', 1))
    assert stored_empty_type(root, "urn:sheet") == expected


def test_actual_empty_xml_cell_string_gap_whitespace_and_formula_remain_distinct():
    w = Workbook()
    s = w.active
    s.append(["ID", "Value"])
    for i, value in enumerate([None, None, "", " ", "=1+1"], 1):
        s.append([f"{i:04d}", value])
    s["B2"].number_format = "@"  # Explicit XML cell; B3 has no stored cell element.
    data = io.BytesIO()
    w.save(data)
    w.close()
    doc = observe(data.getvalue(), "xlsx")
    cells = {
        n["semantic"]["cell"]["coordinate"]: ref
        for ref, n in doc.nodes.items()
        if n.get("semantic", {}).get("cell", {}).get("coordinate")
    }
    assert "B2" in cells and "B3" not in cells and "B4" in cells
    values = {c: doc.nodes[ref]["semantic"]["value"] for c, ref in cells.items()}
    assert values["B2"] == {
        "kind": "blank",
        "raw": "",
        "rawType": "n",
        "observation": "stored_empty_cell",
    }
    assert values["B4"]["kind"] == "string" and values["B4"]["value"] == ""
    assert values["B5"]["value"] == " "
    assert values["B6"]["kind"] == "formula" and values["B6"]["formula"] == "=1+1"
    assert not values["B6"]["cachedAvailable"] and values["B6"]["cachedValue"] is None
    for coordinate in ["B2", "B4"]:
        bid = preferred_binding(doc.bindings, cells[coordinate])
        value, raw, binding = _read(doc.bindings[bid], "string", doc.nodes)
        assert value == raw == "" and binding.path == "/semantic/value/raw"
        assert binding.sourceRef == cells[coordinate]


@pytest.mark.parametrize("shared_empty_string", [False, True])
def test_styled_merge_coverage_does_not_create_independent_blank_cells(shared_empty_string):
    w = Workbook()
    s = w.active
    s["A1"] = "Group"
    s["A1"].border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    s.merge_cells("A1:C2")
    s["A3"] = "Item"
    s["B3"] = "Value"
    s["C3"].number_format = "@"
    data = io.BytesIO()
    w.save(data)
    w.close()
    if shared_empty_string:
        # Read-only OpenPyXL yields '' rather than None for an empty shared string.
        output = io.BytesIO()
        with zipfile.ZipFile(data) as source, zipfile.ZipFile(output, "w") as target:
            for item in source.infolist():
                content = source.read(item)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(content)
                    cell = next(
                        c for c in root.iter() if c.tag.endswith("}c") and c.get("r") == "C1"
                    )
                    cell.set("t", "s")
                    ET.SubElement(
                        cell, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v"
                    ).text = "0"
                    content = ET.tostring(root)
                elif item.filename == "[Content_Types].xml":
                    root = ET.fromstring(content)
                    ET.SubElement(
                        root,
                        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                        {
                            "PartName": "/xl/sharedStrings.xml",
                            "ContentType": (
                                "application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sharedStrings+xml"
                            ),
                        },
                    )
                    content = ET.tostring(root)
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    root = ET.fromstring(content)
                    ET.SubElement(
                        root,
                        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                        {
                            "Id": "rIdEmptyString",
                            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                            "Target": "sharedStrings.xml",
                        },
                    )
                    content = ET.tostring(root)
                target.writestr(item, content)
            target.writestr(
                "xl/sharedStrings.xml",
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t></t></si></sst>',
            )
        data = output
    doc = observe(data.getvalue(), "xlsx")
    table = next(iter(doc.tables.values()))
    cells = {(c["row"], c["col"]): c for c in table["cells"]}
    assert set(cells) == {(0, 0), (2, 0), (2, 1), (2, 2)}
    assert cells[(0, 0)]["rowSpan"] == 2 and cells[(0, 0)]["colSpan"] == 3


def test_sparse_merge_index_matches_rectangles_without_expanding_covered_cells():
    rng = random.Random(47)
    ranges = [
        {
            "origin": {"row": rng.randrange(1, 50), "col": rng.randrange(1, 50)},
            "row_span": rng.randrange(1, 12),
            "col_span": rng.randrange(1, 12),
        }
        for _ in range(200)
    ]
    index = MergeCoverage(ranges)
    for _ in range(1000):
        row, col = rng.randrange(1, 65), rng.randrange(1, 65)
        expected = any(
            (row, col) != (r["origin"]["row"], r["origin"]["col"])
            and r["origin"]["row"] <= row < r["origin"]["row"] + r["row_span"]
            and r["origin"]["col"] <= col < r["origin"]["col"] + r["col_span"]
            for r in ranges
        )
        assert index.covered(row, col) == expected
    huge = MergeCoverage([{"origin": {"row": 1, "col": 1}, "row_span": 1000000, "col_span": 16000}])
    assert not huge.covered(1, 1) and huge.covered(1000000, 16000)
    assert not huge.covered(1000001, 1)
