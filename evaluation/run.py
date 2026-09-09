#!/usr/bin/env python3
"""Run explicitly configured real-model qualification on public synthetic documents.

This produces evidence for review, not an automatic semantic-accuracy certificate.
No model endpoint, credentials, or private document paths enter the evidence.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from document_files import __version__
from document_files.api import ChatCompletionsClient, extract_schema
from document_files.engine import create_hwpx

ROOT = Path(__file__).resolve().parents[1]
LINES = [
    "Reference measurements / 참조 측정 기록",
    "All lengths are in mm. The note applies to sample B only.",
    "Sample A: length 0012.40; count 0; enabled false.",
    "Sample B: length 15.60; count blank; enabled true.",
    "Note B: exclude from the total if the inspection status is pending.",
    "Inspection status B: pending. No status is supplied for sample A.",
]


def fixtures(root: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook

    text = "\n".join(LINES) + "\n"
    for suffix in ("txt", "md"):
        (root / f"measurements.{suffix}").write_text(text, encoding="utf-8")
    (root / "measurements.html").write_text(
        "<!doctype html><html><body>"
        + "".join(f"<p>{line}</p>" for line in LINES)
        + "</body></html>",
        encoding="utf-8",
    )
    doc = Document()
    for line in LINES:
        doc.add_paragraph(line)
    doc.save(root / "measurements.docx")
    book = Workbook()
    sheet = book.active
    for line in LINES:
        sheet.append([line])
    sheet.append(["Native decimal", 12.4, "=B7*2"])
    sheet["B7"].number_format = "0000.00"
    book.save(root / "measurements.xlsx")
    create_hwpx(
        root / "measurements.hwpx",
        plan={
            "schemaVersion": "hwpx.document_plan.v1",
            "title": "Reference measurements",
            "blocks": [{"type": "paragraph", "text": line} for line in LINES],
        },
    )
    return sorted(root.glob("measurements.*"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--kind", required=True, choices=("local", "cloud"))
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--config", type=Path, help="Explicit administrator runtime profile file")
    parser.add_argument("--profile", help="Local-pack or configured cloud profile name")
    parser.add_argument("--options", type=Path, help="Explicit extraction budgets/options JSON")
    parser.add_argument(
        "--case-format", action="append", help="Run only affected formats; repeat as needed"
    )
    parser.add_argument("--holdout-only", action="store_true")
    parser.add_argument(
        "--holdout",
        type=Path,
        help="Independent public case manifest; never use private work documents",
    )
    parser.add_argument("--artifact-inventory", type=Path)
    parser.add_argument(
        "--core-artifact", help="Selected executed document-files wheel artifact ID"
    )
    parser.add_argument(
        "--evaluator-artifact", help="Selected evaluator source distribution artifact ID"
    )
    parser.add_argument(
        "--artifact", action="append", default=[], help="Exact inventory artifact ID"
    )
    parser.add_argument("--evidence-root", type=Path, help="Root for inventory references")
    args = parser.parse_args()
    if bool(args.config) != bool(args.profile):
        parser.error("--config and --profile must be supplied together")
    if args.holdout_only and not args.holdout:
        parser.error("--holdout-only requires --holdout")
    if args.config:
        from document_files.profiles import profile_clients

        with profile_clients(args.config, args.profile) as (client, observation):
            return evaluate(args, client, observation)
    # Stop before producing evidence if not explicitly configured.
    return evaluate(args, ChatCompletionsClient.from_environment(), None)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path: Path):
    def unique_keys(pairs):
        value = {}
        for key, item in pairs:
            if key in value:
                raise ValueError("Duplicate JSON key in evaluation evidence")
            value[key] = item
        return value

    return json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=unique_keys)


def write_json(path: Path, value) -> None:
    # Evidence is append-only; even a review must use a new file.
    with path.open("x", encoding="utf-8") as stream:
        stream.write(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def checked_reference(root: Path, reference: dict) -> Path:
    if not isinstance(reference, dict) or set(reference) != {"path", "sha256"}:
        raise ValueError("Expected a path and SHA256 reference")
    relative = Path(reference["path"])
    path = root / relative
    if (
        relative.is_absolute()
        or ".." in relative.parts
        or not relative.parts
        or not path.resolve().is_relative_to(root.resolve())
        or any(part.is_symlink() for part in (path, *path.parents))
    ):
        raise ValueError("Unsafe evidence reference")
    if sha256(path) != reference["sha256"]:
        raise ValueError("Evidence checksum mismatch")
    return path


def read_specification(path: Path) -> dict:
    spec = read_json(path)
    if (
        not isinstance(spec, dict)
        or spec.get("schemaVersion") != "document-files.review-specification.v1"
    ):
        raise ValueError("Unknown review specification version")
    criteria = spec.get("criteria")
    if not isinstance(criteria, list) or not criteria:
        raise ValueError("Review specification needs pre-inference criteria")
    identifiers = set()
    for criterion in criteria:
        if (
            not isinstance(criterion, dict)
            or not isinstance(criterion.get("id"), str)
            or not re.fullmatch(r"[A-Za-z0-9_-]+", criterion["id"])
            or criterion["id"] in identifiers
            or not isinstance(criterion.get("description"), str)
            or not criterion["description"].strip()
            or "expected" not in criterion
        ):
            raise ValueError("Invalid or duplicate review criterion")
        identifiers.add(criterion["id"])
    return spec


def load_holdouts(manifest_path: Path) -> list[dict]:
    manifest = read_json(manifest_path)
    if (
        not isinstance(manifest, dict)
        or manifest.get("schemaVersion") != "document-files.evaluation-cases.v2"
    ):
        raise ValueError("Holdouts require evaluation-cases.v2 metadata and frozen specifications")
    if manifest.get("publicData") is not True:
        raise ValueError("Holdout cases must be explicitly declared public data")
    if not isinstance(manifest.get("cases"), list) or not manifest["cases"]:
        raise ValueError("Holdout manifest requires a nonempty case list")
    cases, identifiers = [], set()
    for item in manifest["cases"]:
        allowed = {
            "id",
            "path",
            "languages",
            "printed",
            "pdfKind",
            "longDocument",
            "pageCount",
            "regionCount",
            "reviewSpecification",
            "recognitionCostReceipt",
        }
        if not isinstance(item, dict) or set(item) - allowed:
            raise ValueError("Unknown holdout metadata; put expectations only in the specification")
        identifier = item.get("id")
        if (
            not isinstance(identifier, str)
            or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,99}", identifier)
            or identifier in identifiers
        ):
            raise ValueError("Case IDs must be unique safe identifiers")
        identifiers.add(identifier)
        path = (manifest_path.parent / item["path"]).resolve()
        if path.suffix.lower().lstrip(".") not in {
            "txt",
            "md",
            "html",
            "docx",
            "hwp",
            "hwpx",
            "xlsx",
            "pptx",
            "pdf",
        }:
            raise ValueError("Unsupported evaluation input format")
        languages = item.get("languages")
        if (
            not isinstance(languages, list)
            or not languages
            or any(not isinstance(value, str) or not value.strip() for value in languages)
            or type(item.get("printed")) is not bool
            or type(item.get("longDocument")) is not bool
            or item.get("pdfKind") not in ("text", "scan", "mixed", "not-applicable")
        ):
            raise ValueError("Case language, printed, PDF kind and long-document metadata required")
        if (path.suffix.lower() == ".pdf") == (item["pdfKind"] == "not-applicable"):
            raise ValueError("PDF kind must agree with the input format")
        for count in ("pageCount", "regionCount"):
            if count in item and (type(item[count]) is not int or item[count] < 1):
                raise ValueError("Page and region counts must be positive integers")
        if item["longDocument"] and any(
            item.get(key, 0) < 3 for key in ("pageCount", "regionCount")
        ):
            raise ValueError("Long-document cases require at least three pages and regions")
        spec_path = checked_reference(manifest_path.parent, item["reviewSpecification"])
        read_specification(spec_path)
        if "recognitionCostReceipt" in item:
            checked_reference(manifest_path.parent, item["recognitionCostReceipt"])
        cases.append({**item, "path": path, "holdout": True, "manifestRoot": manifest_path.parent})
    return cases


def case_options(case: dict, overrides: dict) -> dict:
    from document_files.api import ExtractionOptions

    # Expectations never enter options or the extractor input. Overrides are operator-only.
    return ExtractionOptions.model_validate(
        {
            "reconstructionContext": False,
            "maxModelCalls": 64 if case["longDocument"] else 12,
            "completionSeconds": 3600 if case["longDocument"] else 900,
            **overrides,
        }
    ).model_dump()


def external_reference(args, path: Path) -> dict:
    root = getattr(args, "evidence_root", None)
    if root is None:
        raise ValueError("Artifact references require --evidence-root")
    root, path = root.expanduser().resolve(), path.expanduser().absolute()
    try:
        relative = path.relative_to(root)
    except ValueError as exc:
        raise ValueError("Referenced evidence must be inside --evidence-root") from exc
    reference = {"path": relative.as_posix(), "sha256": sha256(path)}
    checked_reference(root, reference)
    return reference


def verify_executed_source(args, inventory, client, observation):
    import document_files

    spec = importlib.util.spec_from_file_location(
        "document_files_qualification_identity", Path(__file__).with_name("execution_identity.py")
    )
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    return helper.verify_execution(
        inventory,
        args.artifact_inventory,
        args.evidence_root,
        args.artifact,
        getattr(args, "core_artifact", None),
        getattr(args, "evaluator_artifact", None),
        Path(document_files.__file__).resolve().parent,
        Path(__file__).resolve().parent,
        client.identity,
        observation.identity if observation else None,
    )


def evaluate(args, client, observation):
    started = time.monotonic()
    overrides = read_json(args.options) if args.options else {}
    if not isinstance(overrides, dict):
        raise ValueError("Options must be an object")
    output = args.output.expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise ValueError("Use a new empty evidence directory; existing results are not overwritten")
    source_commit, dirty, core_execution = None, None, None
    inventory_path = getattr(args, "artifact_inventory", None)
    artifacts = getattr(args, "artifact", [])
    if bool(inventory_path) != bool(artifacts) or len(set(artifacts)) != len(artifacts):
        raise ValueError("Supply both --artifact-inventory and unique --artifact IDs")
    inventory_reference = external_reference(args, inventory_path) if inventory_path else None
    if inventory_path:
        inventory = read_json(inventory_path)
        if (
            inventory.get("schemaVersion") != "document-files.artifact-inventory.v2"
            or inventory.get("version") != __version__
            or inventory.get("dirtySource") is not False
        ):
            raise ValueError("Inventory must identify a clean source and matching product version")
        available = {item["id"] for item in inventory["artifacts"]}
        if len(available) != len(inventory["artifacts"]):
            raise ValueError("Inventory artifact IDs must be unique")
        if not set(artifacts) <= available:
            raise ValueError("Requested artifact is not in the inventory")
        for item in inventory["artifacts"]:
            if item["id"] in artifacts:
                checked_reference(inventory_path.parent, {k: item[k] for k in ("path", "sha256")})
        core_execution, source_commit = verify_executed_source(args, inventory, client, observation)
        dirty = False
    else:
        # Development-only identity. A source archive/untrusted-owner Git failure stays
        # explicitly unknown, never a clean commit and never a reason to change safe.directory.
        try:
            source_commit = subprocess.check_output(
                ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True, stderr=subprocess.DEVNULL
            ).strip()
            dirty = bool(
                subprocess.check_output(
                    ["git", "status", "--porcelain"], cwd=ROOT, stderr=subprocess.DEVNULL
                )
            )
        except (OSError, subprocess.CalledProcessError):
            pass
    records = []
    with tempfile.TemporaryDirectory(prefix="document-files-qualification-") as directory:
        inputs = [] if args.holdout_only else fixtures(Path(directory))
        cases = [
            {
                "path": path,
                "id": "development-" + path.suffix[1:],
                "holdout": False,
                "languages": ["en", "ko"],
                "printed": True,
                "pdfKind": "not-applicable",
                "longDocument": False,
            }
            for path in inputs
        ]
        if args.holdout:
            cases.extend(load_holdouts(args.holdout))
        if args.case_format:
            cases = [c for c in cases if c["path"].suffix[1:].lower() in args.case_format]
        if not cases:
            raise ValueError("No cases match the explicit selection")
        if len({case["id"] for case in cases}) != len(cases):
            raise ValueError("Case IDs must not collide with development case IDs")
        # Freeze every input, expected specification and option before any inference.
        # The model sees only the input snapshot and validated product options.
        (output / "inputs").mkdir()
        (output / "review-specifications").mkdir()
        (output / "recognition-receipts").mkdir()
        prepared = []
        for case in cases:
            record = {
                key: value
                for key, value in case.items()
                if key
                not in {
                    "path",
                    "manifestRoot",
                    "reviewSpecification",
                    "recognitionCostReceipt",
                }
            }
            record["format"] = case["path"].suffix.lstrip(".").lower()
            input_name = f"inputs/{case['id']}.{record['format']}"
            shutil.copyfile(case["path"], output / input_name)
            record.update(inputPath=input_name, inputSha256=sha256(output / input_name))
            for key, folder in (
                ("reviewSpecification", "review-specifications"),
                ("recognitionCostReceipt", "recognition-receipts"),
            ):
                if key in case:
                    original = checked_reference(case["manifestRoot"], case[key])
                    name = f"{folder}/{case['id']}.json"
                    shutil.copyfile(original, output / name)
                    record[key] = {"path": name, "sha256": sha256(output / name)}
            record["options"] = case_options(case, overrides)
            prepared.append(record)
        for record in prepared:
            input_path = output / record["inputPath"]
            inference_started = time.monotonic()
            result = extract_schema(
                str(input_path),
                model_client=client,
                retain=False,
                options=record["options"],
                observation_backend=observation,
            )
            elapsed = time.monotonic() - inference_started
            if sha256(input_path) != record["inputSha256"]:
                raise ValueError("Input changed during extraction")
            if "reviewSpecification" in record:
                checked_reference(output, record["reviewSpecification"])
            result_name = f"{args.kind}-{record['id']}.json"
            write_json(output / result_name, result)
            record.update(
                resultPath=result_name,
                resultSha256=sha256(output / result_name),
                extractionStatus=result["extraction"]["status"],
                structurallyValid=result["validation"]["valid"],
                elapsedSeconds=elapsed,
                usage=result["extraction"].get("usage"),
                recognitionCostStatus="receipt-provided"
                if "recognitionCostReceipt" in record
                else "not-separately-measured",
                semanticReview={"status": "pending"},
            )
            records.append(record)
    if inventory_path:
        after, after_commit = verify_executed_source(args, inventory, client, observation)
        if after != core_execution or after_commit != source_commit:
            raise ValueError("Executed core/evaluator changed during qualification")
    report = {
        "schemaVersion": "document-files.model-qualification.v2",
        "version": __version__,
        "sourceCommit": source_commit,
        "dirtySource": dirty,
        "endpointKind": args.kind,
        "executionKind": "actual-model",
        "model": client.identity,
        "observation": observation.identity if observation else None,
        "optionOverrides": overrides,
        "elapsedSeconds": time.monotonic() - started,
        "passed": False,
        "cases": records,
        "artifacts": artifacts,
        "reason": "Independent semantic review and release coverage remain pending.",
    }
    if core_execution:
        report["coreExecution"] = core_execution
    if inventory_reference:
        report["artifactInventory"] = inventory_reference
    if os.environ.get("DOCUMENT_FILES_EXECUTION_RUN_ID"):
        report["executionRunId"] = os.environ["DOCUMENT_FILES_EXECUTION_RUN_ID"]
    name = output / f"{args.kind}_model.json"
    write_json(name, report)
    print(f"Evidence written: {name}; semantic qualification remains pending")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
